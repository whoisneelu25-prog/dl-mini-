import pytest
import sys
import os
import json
import pptx
from pptx.util import Inches

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from utils import export_to_pptx, export_to_markdown, export_to_json, get_system_info


@pytest.fixture
def sample_presentation():
    return {
        "id": "pres_test_123",
        "topic": "Artificial Intelligence in Healthcare",
        "objective": "Explain how AI is transforming healthcare diagnosis and patient care",
        "audience": "College Students",
        "difficulty": "Intermediate",
        "presentation_type": "Technical Seminar",
        "num_slides": 3,
        "duration": 6,
        "domain": "Healthcare",
        "points_refined": 1,
        "slides": [
            {
                "slide_number": 1,
                "title": "Introduction to AI in Healthcare",
                "purpose": "Introduce clinical AI paradigms",
                "time_minutes": 2.0,
                "bullets": [
                    "Overview of AI and machine learning in modern medicine.",
                    "Integration with electronic health record workflows.",
                    "Enhancement of diagnostic accuracy and clinical throughput.",
                ],
                "case_study": "Deep learning models reducing false-negative mammogram screening rates by 15%.",
                "speaker_notes": "Introduce the motivation for healthcare AI to student audience.",
            },
            {
                "slide_number": 2,
                "title": "AI Technologies in Healthcare",
                "purpose": "Analyze core convolutional and transformer models",
                "time_minutes": 2.0,
                "bullets": [
                    "Convolutional Neural Networks for volumetric MRI analysis.",
                    "Natural Language Processing for unstructured clinical note extraction.",
                    "Federated learning preserving patient medical privacy across hospitals.",
                ],
                "case_study": "Federated hospital consortium benchmarking rare oncology detection.",
                "speaker_notes": "Highlight privacy constraints and convolutional architectures.",
            },
            {
                "slide_number": 3,
                "title": "Conclusion & Future Horizons",
                "purpose": "Summarize clinical impact and open frontiers",
                "time_minutes": 2.0,
                "bullets": [
                    "Transformative synthesis of clinician expertise and machine predictions.",
                    "Regulatory pathways including FDA SaMD certification.",
                    "Future frontiers in personalized genomic medicine.",
                ],
                "case_study": "Real-time clinical decision support deployed in intensive care units.",
                "speaker_notes": "Conclude presentation with emphasis on human-in-the-loop clinical validation.",
            },
        ],
    }


def test_pptx_export(sample_presentation, tmp_path):
    out_file = str(tmp_path / "test_presentation.pptx")
    res_path = export_to_pptx(sample_presentation, out_file)

    assert os.path.exists(res_path)
    assert os.path.getsize(res_path) > 1000

    # Inspect with python-pptx
    prs = pptx.Presentation(res_path)
    # Title slide + 3 content slides = 4 total slides
    assert len(prs.slides) == 4

    # Check 16:9 widescreen dimensions
    assert pytest.approx(prs.slide_width.inches, abs=0.01) == 13.333
    assert pytest.approx(prs.slide_height.inches, abs=0.01) == 7.5

    # Check that speaker notes exist on content slides
    slide1 = prs.slides[1]
    assert slide1.has_notes_slide
    notes = slide1.notes_slide.notes_text_frame.text
    assert "Speaker Notes:" in notes or "Allocated Pacing" in notes


def test_markdown_export(sample_presentation, tmp_path):
    out_file = str(tmp_path / "test_presentation.md")
    res_path = export_to_markdown(sample_presentation, out_file)

    assert os.path.exists(res_path)
    with open(res_path, "r", encoding="utf-8") as f:
        content = f.read()

    assert "# Artificial Intelligence in Healthcare" in content
    assert "Slide 01: Introduction to AI in Healthcare" in content
    assert "Allocated Time" in content
    assert "Case Study" in content
    assert "Speaker Notes" in content


def test_json_export(sample_presentation, tmp_path):
    out_file = str(tmp_path / "test_presentation.json")
    res_path = export_to_json(sample_presentation, out_file)

    assert os.path.exists(res_path)
    with open(res_path, "r", encoding="utf-8") as f:
        loaded = json.load(f)

    assert loaded["topic"] == sample_presentation["topic"]
    assert len(loaded["slides"]) == 3
    assert loaded["slides"][0]["title"] == "Introduction to AI in Healthcare"


def test_system_info_structure():
    info = get_system_info()
    assert info["generation_model"] == "google/flan-t5-base"
    assert info["semantic_model"] == "all-MiniLM-L6-v2"
    assert info["embedding_dimension"] == 384
    assert info["similarity_threshold"] == 0.88
    assert info["execution_device"] in {"MPS", "CUDA", "CPU"}
